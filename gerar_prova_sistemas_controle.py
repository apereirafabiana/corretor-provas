# -*- coding: utf-8 -*-
"""Gera provas, folha de respostas e gabaritos de Instrumentação Industrial."""

from __future__ import annotations

import json
import random
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from openpyxl import Workbook
from PIL import Image, ImageOps
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "saida_prova_sistemas_controle"
LOGO = ROOT / "cefet_mg_logo.png"

VERSIONS = ["A", "B", "C", "D"]
TOTAL_POINTS = 10.0
POINTS_PER_QUESTION = 0.50
RANDOM_SEED = 20260428
EXAM_TITLE = "Prova de Sistemas de Controle"
ANSWER_SHEET_TITLE = "Folha de Respostas - Sistemas de Controle"
COURSE_LINE = "Curso Técnico em Eletrônica - Disciplina: Sistemas de Controle"
FOOTER_TITLE = "Sistemas de Controle - Curso Técnico em Eletrônica"

PAGE_W, PAGE_H = A4
BLACK = (0.0, 0.0, 0.0)
DARK_GRAY = (0.20, 0.20, 0.20)
MID_GRAY = (0.55, 0.55, 0.55)
LIGHT_GRAY = (0.985, 0.985, 0.985)
TEXT_GRAY = BLACK
BLUE = BLACK
GREEN = DARK_GRAY


OLD_QUESTIONS = [
    {
        "text": "Em instrumentação industrial, o termo processo pode ser definido como:",
        "options": {
            "A": "a variável que indica diretamente o estado desejado do produto.",
            "B": "a energia ou o material no qual a variável é controlada.",
            "C": "a variável sobre a qual o controlador atua para corrigir o sistema.",
            "D": "uma operação ou série de operações realizada em determinado equipamento, na qual varia pelo menos uma característica física ou química de um material.",
            "E": "o valor desejado para uma variável controlada.",
        },
        "answer": "D",
    },
    {
        "text": "Em uma planta industrial, são exemplos de variáveis de processo:",
        "options": {
            "A": "válvula, bomba, motor e CLP.",
            "B": "pressão, temperatura, vazão, nível e pH.",
            "C": "controlador, transmissor, indicador e registrador.",
            "D": "tubulação, flange, tanque e serpentina.",
            "E": "alarme, intertravamento, botoeira e relé.",
        },
        "answer": "B",
    },
    {
        "text": "Em um sistema de aquecimento de água por vapor, deseja-se manter constante a temperatura da água de saída. O controlador atua sobre uma válvula que regula a vazão de vapor.\n\nNesse caso, a variável controlada e a variável manipulada são, respectivamente:",
        "options": {
            "A": "vazão de vapor e temperatura da água de saída.",
            "B": "pressão do vapor e abertura da válvula.",
            "C": "temperatura da água de saída e vazão de vapor.",
            "D": "temperatura ambiente e vazão de condensado.",
            "E": "nível do tanque e temperatura do vapor.",
        },
        "answer": "C",
    },
    {
        "text": "Uma malha de controle é classificada como malha fechada quando:",
        "options": {
            "A": "o controlador atua sem receber informação da variável controlada.",
            "B": "a saída do processo não interfere na ação de controle.",
            "C": "a variável controlada é medida, comparada ao valor desejado e usada para corrigir o processo.",
            "D": "o operador ajusta manualmente a válvula sem instrumento de medição.",
            "E": "não existe elemento sensor no processo.",
        },
        "answer": "C",
    },
    {
        "text": "Em uma malha aberta de controle de temperatura:",
        "options": {
            "A": "a temperatura medida é sempre comparada com o set point.",
            "B": "a informação da variável controlada não é utilizada para corrigir a entrada do sistema.",
            "C": "o controlador corrige automaticamente qualquer perturbação externa.",
            "D": "existe obrigatoriamente realimentação negativa.",
            "E": "o sensor envia sinal diretamente ao elemento final de controle.",
        },
        "answer": "B",
    },
    {
        "text": "Os sistemas de controle automático são compostos, basicamente, por:",
        "options": {
            "A": "unidade de medida, unidade de controle e elemento final de controle.",
            "B": "compressor, válvula manual e tubulação.",
            "C": "painel elétrico, disjuntor e fusível.",
            "D": "reservatório, motor e operador.",
            "E": "bomba, sensor e fonte de alimentação apenas.",
        },
        "answer": "A",
    },
    {
        "text": "Um instrumento que detecta variações na variável medida/controlada por meio de um elemento primário e transmite essa informação à distância é denominado:",
        "options": {
            "A": "indicador.",
            "B": "transmissor.",
            "C": "elemento final de controle.",
            "D": "registrador.",
            "E": "valvula de controle.",
        },
        "answer": "B",
    },
    {
        "text": "Um controlador industrial tem como função principal:",
        "options": {
            "A": "modificar diretamente a quantidade de material ou energia do processo.",
            "B": "apenas indicar localmente o valor da variável.",
            "C": "comparar a variável de processo com o valor desejado e fornecer um sinal de correção.",
            "D": "registrar graficamente os valores medidos ao longo do tempo.",
            "E": "converter um sinal pneumático em sinal hidráulico.",
        },
        "answer": "C",
    },
    {
        "text": "Uma válvula de controle, um inversor de frequência e uma válvula solenoide são exemplos de:",
        "options": {
            "A": "elementos primários de medição.",
            "B": "instrumentos indicadores.",
            "C": "elementos finais de controle.",
            "D": "registradores gráficos.",
            "E": "sensores passivos.",
        },
        "answer": "C",
    },
    {
        "text": "Um transmissor de pressão possui faixa nominal de medição de 100 kPa a 500 kPa. O span desse instrumento é:",
        "options": {
            "A": "100 kPa.",
            "B": "400 kPa.",
            "C": "500 kPa.",
            "D": "600 kPa.",
            "E": "300 kPa.",
        },
        "answer": "B",
    },
    {
        "text": "Um transmissor possui entrada de 0 a 200 °C e saída de 4 mA a 20 mA. Considerando comportamento linear, quando a temperatura estiver em 50% da faixa, a saída será:",
        "options": {
            "A": "4 mA.",
            "B": "8 mA.",
            "C": "10 mA.",
            "D": "12 mA.",
            "E": "20 mA.",
        },
        "answer": "D",
    },
    {
        "text": "Um transmissor de nível possui saída de 4 mA a 20 mA. Durante uma inspeção, o técnico mede uma corrente de 16 mA. Considerando relação linear, esse valor corresponde a:",
        "options": {
            "A": "25% da faixa.",
            "B": "50% da faixa.",
            "C": "60% da faixa.",
            "D": "75% da faixa.",
            "E": "100% da faixa.",
        },
        "answer": "D",
    },
    {
        "text": "Um instrumento possui range de 50 °C a 150 °C e exatidão de ±0,5% do span. Ao indicar 80 °C, o valor verdadeiro da temperatura estará no intervalo:",
        "options": {
            "A": "79,5 °C a 80,5 °C.",
            "B": "79,6 °C a 80,4 °C.",
            "C": "79,0 °C a 81,0 °C.",
            "D": "75,0 °C a 85,0 °C.",
            "E": "50,5 °C a 150,5 °C.",
        },
        "answer": "A",
    },
    {
        "text": "Um instrumento possui range de 0 °C a 200 °C e zona morta de ±0,1% do span. A zona morta desse instrumento é:",
        "options": {
            "A": "±0,02 °C.",
            "B": "±0,1 °C.",
            "C": "±0,2 °C.",
            "D": "±1,0 °C.",
            "E": "±2,0 °C.",
        },
        "answer": "C",
    },
    {
        "text": "Durante a calibração de um transmissor, observa-se que, para o mesmo valor de entrada, a saída do instrumento apresenta valores diferentes quando a variável está aumentando e quando está diminuindo.\n\nEssa característica é denominada:",
        "options": {
            "A": "resolução.",
            "B": "sensibilidade.",
            "C": "histerese.",
            "D": "fundo de escala.",
            "E": "linearidade.",
        },
        "answer": "C",
    },
    {
        "text": "A sensibilidade de um instrumento está relacionada:",
        "options": {
            "A": "à menor variação perceptível no mostrador.",
            "B": "à diferença entre o maior e o menor valor da escala.",
            "C": "à razão entre a variação do sinal de saída e a correspondente variação da grandeza de entrada.",
            "D": "ao erro máximo admissível do instrumento.",
            "E": "à diferença entre o valor indicado e o valor real.",
        },
        "answer": "C",
    },
    {
        "text": "Segundo a identificação funcional da ISA S5.1, a tag TIC-103 representa, de modo geral:",
        "options": {
            "A": "transmissor indicador de corrente da malha 103.",
            "B": "controlador indicador de temperatura da malha 103.",
            "C": "transmissor de nível da malha 103.",
            "D": "chave indicadora de pressão da malha 103.",
            "E": "registrador de vazão da malha 103.",
        },
        "answer": "B",
    },
    {
        "text": "De acordo com a ISA S5.1, um transmissor de pressão diferencial utilizado funcionalmente para medição de nível deve ser identificado como:",
        "options": {
            "A": "PDT, pois o instrumento mede pressão diferencial internamente.",
            "B": "PT, pois a grandeza física original é pressão.",
            "C": "LT, pois a função desempenhada na malha é transmissão de nível.",
            "D": "LI, pois qualquer instrumento de nível é apenas indicador.",
            "E": "LSH, pois todo instrumento de nível atua como chave de nível alto.",
        },
        "answer": "C",
    },
    {
        "text": "Em uma tubulação com fluido quente, corrosivo e em alta velocidade, deseja-se instalar um sensor de temperatura intrusivo, permitindo manutenção sem abertura do processo. A solução mais adequada é:",
        "options": {
            "A": "instalar o sensor diretamente no fluido, sem proteção.",
            "B": "utilizar cabo de compensação.",
            "C": "utilizar termopoço.",
            "D": "utilizar apenas indicador local.",
            "E": "substituir o sensor por uma valvula manual.",
        },
        "answer": "C",
    },
    {
        "text": "Sobre sensores de temperatura industriais, assinale a alternativa correta.",
        "options": {
            "A": "Termopares medem diretamente temperatura absoluta e dispensam compensação da junta fria.",
            "B": "Termistores NTC aumentam sua resistência elétrica com o aumento da temperatura.",
            "C": "RTDs, como o Pt100, medem temperatura com base na variação da resistência elétrica de um metal.",
            "D": "Sensores bimetálicos funcionam com base no efeito Seebeck.",
            "E": "Termômetros de dilatação de líquido são os mais indicados para transmissão direta de sinal 4 mA a 20 mA.",
        },
        "answer": "C",
    },
]


QUESTIONS = [
    {
        "part": "Parte I - Instrumentação Geral",
        "text": "Em uma unidade industrial, uma malha de controle automático tem como finalidade manter determinada variável em um valor desejado, mesmo diante de perturbações no processo. Considerando os conceitos básicos de instrumentação, uma malha de controle pode ser definida como:",
        "options": {
            "A": "o conjunto de tubulações e válvulas manuais de um processo industrial.",
            "B": "a sequência de operações químicas realizadas em um equipamento.",
            "C": "a interconexão de dispositivos com o objetivo de medir, controlar e atuar sobre uma variável.",
            "D": "o valor máximo que um instrumento pode medir sem sofrer dano permanente.",
            "E": "a diferença entre o valor medido e o valor verdadeiro da variável.",
        },
        "answer": "C",
    },
    {
        "text": "Em um sistema de aquecimento de água por vapor, a temperatura da água de saída deve ser mantida em 85 °C. Um transmissor mede a temperatura da água, o controlador compara esse valor com o set point e atua sobre uma válvula que regula a vazão de vapor.\n\nNesse sistema, a variável controlada e a variável manipulada são, respectivamente:",
        "options": {
            "A": "temperatura da água de saída e vazão de vapor.",
            "B": "vazão de vapor e temperatura da água de saída.",
            "C": "pressão do vapor e vazão da água de entrada.",
            "D": "abertura da válvula e temperatura ambiente.",
            "E": "nível da água e pressão atmosférica.",
        },
        "answer": "A",
    },
    {
        "text": "Um técnico observa que determinado sistema industrial aciona uma resistência elétrica durante 10 minutos a cada hora, independentemente da temperatura real do equipamento. Não há sensor nem comparação entre temperatura medida e valor desejado.\n\nEsse sistema caracteriza uma:",
        "options": {
            "A": "malha fechada, pois existe atuação sobre o processo.",
            "B": "malha fechada, pois há correção automática da variável controlada.",
            "C": "malha aberta, pois a variável controlada não é usada para corrigir a entrada.",
            "D": "malha aberta, pois obrigatoriamente há realimentação negativa.",
            "E": "malha fechada com controlador proporcional.",
        },
        "answer": "C",
    },
    {
        "text": "Um transmissor de pressão possui faixa de medição ajustada de 2 bar a 10 bar e saída padronizada de 4 mA a 20 mA. A pressão do processo está em 6 bar. Admitindo resposta linear, a corrente de saída será:",
        "options": {
            "A": "8 mA.",
            "B": "10 mA.",
            "C": "12 mA.",
            "D": "14 mA.",
            "E": "16 mA.",
        },
        "answer": "C",
    },
    {
        "text": "Um transmissor de nível possui range de 0 a 600 mm e saída de 4 mA a 20 mA. Durante uma inspeção, foi medida uma corrente de saída de 14,4 mA. Considerando comportamento linear, o nível correspondente é:",
        "options": {
            "A": "240 mm.",
            "B": "300 mm.",
            "C": "360 mm.",
            "D": "390 mm.",
            "E": "420 mm.",
        },
        "answer": "D",
    },
    {
        "text": "Um instrumento de medição industrial possui faixa de 100 V a 300 V, com exatidão de ±0,4% do span. Ao indicar 250 V, o valor verdadeiro da tensão estará compreendido entre:",
        "options": {
            "A": "249,2 V e 250,8 V.",
            "B": "249,4 V e 250,6 V.",
            "C": "249,6 V e 250,4 V.",
            "D": "248,8 V e 251,2 V.",
            "E": "249,0 V e 251,0 V.",
        },
        "answer": "A",
    },
    {
        "text": "Um instrumento possui faixa de medição de -50 °C a 50 °C e zona morta de 1% do span. O valor da zona morta desse instrumento é:",
        "options": {
            "A": "0,1 °C.",
            "B": "0,5 °C.",
            "C": "1,0 °C.",
            "D": "2,0 °C.",
            "E": "5,0 °C.",
        },
        "answer": "C",
    },
    {
        "text": "Um transdutor industrial recebe uma pressão de 0 psi a 100 psi e fornece saída elétrica linear de 1 V a 5 V. A sensibilidade desse transdutor é:",
        "options": {
            "A": "0,01 V/psi.",
            "B": "0,02 V/psi.",
            "C": "0,04 V/psi.",
            "D": "0,05 V/psi.",
            "E": "0,10 V/psi.",
        },
        "answer": "C",
    },
    {
        "text": "Em uma malha de controle de pressão, o controlador eletrônico fornece sinal de saída padronizado de 4 mA a 20 mA. O elemento final de controle é uma válvula pneumática cujo atuador opera com sinal de 3 psi a 15 psi.\n\nDurante a operação, o controlador envia um sinal de 16 mA para posicionar a válvula. Considerando conversão linear, o dispositivo necessário entre o controlador e a válvula e o sinal pneumático correspondente são, respectivamente:",
        "options": {
            "A": "conversor P/I e 9 psi.",
            "B": "conversor I/P e 9 psi.",
            "C": "conversor I/P e 12 psi.",
            "D": "transmissor de pressão e 12 psi.",
            "E": "posicionador eletropneumático e 15 psi.",
        },
        "answer": "C",
    },
    {
        "text": "Segundo a identificação funcional da ISA S5.1, um instrumento construído como transmissor de pressão diferencial, mas utilizado funcionalmente para medir nível em um tanque, deve ser identificado preferencialmente como:",
        "options": {
            "A": "PDT, pois sua construção física é baseada em pressão diferencial.",
            "B": "PT, pois a grandeza primária é pressão.",
            "C": "LT, pois sua função na malha é transmitir nível.",
            "D": "LI, pois todo instrumento de nível deve indicar localmente.",
            "E": "LSH, pois todo instrumento de nível atua como chave de nível alto.",
        },
        "answer": "C",
    },
    {
        "part": "Parte II - Sensores de Temperatura",
        "text": "A medição de temperatura é considerada indireta porque, em geral, o instrumento:",
        "options": {
            "A": "mede sempre a quantidade total de calor armazenada no corpo.",
            "B": "mede exclusivamente a energia cinética das moléculas.",
            "C": "utiliza a alteração de alguma propriedade física do material para inferir a temperatura.",
            "D": "mede apenas a radiação térmica emitida pelo fluido.",
            "E": "mede diretamente o calor sensível e o calor latente do processo.",
        },
        "answer": "C",
    },
    {
        "text": "Em uma tubulação industrial com fluido quente, corrosivo e em alta velocidade, deseja-se instalar um sensor de temperatura intrusivo, mas permitindo sua remoção para manutenção sem abertura da linha de processo.\n\nA solução mais adequada é utilizar:",
        "options": {
            "A": "sensor diretamente imerso no fluido, sem proteção.",
            "B": "termopoço.",
            "C": "cabo de compensação.",
            "D": "junta fria aterrada.",
            "E": "indicador local bimetálico sem proteção.",
        },
        "answer": "B",
    },
    {
        "text": "Em um sistema de proteção térmica de motores elétricos, deseja-se utilizar um sensor cuja resistência aumente significativamente quando a temperatura ultrapassa determinado valor, auxiliando na proteção contra sobreaquecimento.\n\nO sensor mais adequado para essa aplicação é:",
        "options": {
            "A": "termistor NTC.",
            "B": "termistor PTC.",
            "C": "termopar tipo K.",
            "D": "termômetro de dilatação de líquido.",
            "E": "termômetro por pressão de vapor.",
        },
        "answer": "B",
    },
    {
        "text": "Um termômetro por dilatação de líquido utiliza como princípio básico:",
        "options": {
            "A": "a geração de tensão elétrica entre dois metais diferentes.",
            "B": "a variação da resistência elétrica de um semicondutor.",
            "C": "a variação volumétrica de um líquido com a temperatura.",
            "D": "a absorção de radiação infravermelha pelo líquido.",
            "E": "a variação da frequência de vibração de um cristal.",
        },
        "answer": "C",
    },
    {
        "text": "Os termômetros industriais por pressão de vapor, do tipo Bourdon, funcionam porque:",
        "options": {
            "A": "a resistência elétrica do líquido volátil varia linearmente com a temperatura.",
            "B": "o aumento da temperatura altera o equilíbrio líquido-vapor, aumentando a pressão interna do sistema.",
            "C": "dois metais diferentes geram uma força eletromotriz proporcional à temperatura absoluta.",
            "D": "a radiação térmica do fluido deforma diretamente o tubo de Bourdon.",
            "E": "o gás de enchimento sofre ionização quando a temperatura aumenta.",
        },
        "answer": "B",
    },
    {
        "text": "O par bimetálico utilizado em instrumentos de temperatura baseia-se:",
        "options": {
            "A": "na diferença de coeficientes de dilatação térmica entre dois metais unidos.",
            "B": "na variação da resistência elétrica da platina com a temperatura.",
            "C": "no efeito Seebeck entre duas juntas metálicas.",
            "D": "na pressão de vapor de um líquido volátil.",
            "E": "na emissão de radiação eletromagnética por corpos aquecidos.",
        },
        "answer": "A",
    },
    {
        "text": "A respeito dos termopares, assinale a alternativa correta.",
        "options": {
            "A": "Medem diretamente a temperatura absoluta da junta quente.",
            "B": "Funcionam com base no efeito Peltier e dispensam condicionamento de sinal.",
            "C": "Geram uma força eletromotriz da ordem de volts, diretamente aplicada ao CLP.",
            "D": "Medem diferença de temperatura entre junta quente e junta fria, exigindo compensação da junta fria para se obter temperatura absoluta.",
            "E": "São sensores resistivos metálicos de alta linearidade, normalmente baseados em platina.",
        },
        "answer": "D",
    },
    {
        "text": "Em uma instalação industrial, deseja-se medir temperatura com boa precisão e alta estabilidade ao longo do tempo, utilizando um sensor metálico padronizado, bastante empregado em processos industriais.\n\nO sensor mais adequado é:",
        "options": {
            "A": "termistor NTC.",
            "B": "termistor PTC.",
            "C": "Pt100.",
            "D": "termômetro de dilatação de líquido.",
            "E": "par bimetálico simples.",
        },
        "answer": "C",
    },
    {
        "text": "Um sensor Pt100 possui resistência de 100 Ω a 0 °C e coeficiente aproximado de 0,385 Ω/°C. Desprezando não linearidades, sua resistência aproximada a 100 °C será:",
        "options": {
            "A": "100,0 Ω.",
            "B": "103,85 Ω.",
            "C": "119,25 Ω.",
            "D": "138,5 Ω.",
            "E": "185,0 Ω.",
        },
        "answer": "D",
    },
    {
        "text": "Um Pt100 é instalado em ligação a 2 fios. Cada fio de cobre possui 20 m de comprimento e resistência de 0,02 Ω/m. Considerando que a resistência dos fios será somada à resistência do sensor e que o coeficiente do Pt100 é 0,385 Ω/°C, o erro aproximado introduzido pela resistência dos cabos será:",
        "options": {
            "A": "0,52 °C.",
            "B": "1,04 °C.",
            "C": "2,08 °C.",
            "D": "3,85 °C.",
            "E": "5,20 °C.",
        },
        "answer": "C",
    },
]


QUESTIONS = [
    {
        "text": "Em um forno industrial, o controlador compara a temperatura medida com o valor ajustado e envia sinal para a válvula de gás. Nesse caso, a temperatura do forno corresponde a:",
        "options": {
            "A": "variável manipulada do processo.",
            "B": "variável controlada do processo.",
            "C": "saída elétrica do controlador.",
            "D": "perturbação aplicada ao sistema.",
            "E": "elemento final de controle.",
        },
        "answer": "B",
    },
    {
        "text": "Um sistema que aciona uma resistência elétrica por tempo fixo, sem medir a temperatura resultante, caracteriza-se como:",
        "options": {
            "A": "malha fechada com sensor de retorno.",
            "B": "sistema com compensação em tempo real.",
            "C": "malha aberta sujeita a perturbações.",
            "D": "controle realimentado com comparação.",
            "E": "sistema de controle adaptativo.",
        },
        "answer": "C",
    },
    {
        "text": "Na malha de aquecimento de água, o controlador atua sobre a vazão de vapor para manter a temperatura. A vazão de vapor é a:",
        "options": {
            "A": "variável medida.",
            "B": "variável controlada.",
            "C": "variável de referência.",
            "D": "variável manipulada.",
            "E": "grandeza de indicação.",
        },
        "answer": "D",
    },
    {
        "text": "Um transmissor está calibrado para 50 a 250 °C e fornece 4 a 20 mA. Para uma entrada de 150 °C, admitindo linearidade, a saída será:",
        "options": {
            "A": "8 mA.",
            "B": "10 mA.",
            "C": "12 mA.",
            "D": "14 mA.",
            "E": "16 mA.",
        },
        "answer": "C",
    },
    {
        "text": "Um instrumento com faixa de 0 a 200 °C possui zona morta de ±0,1% do span. A variação de temperatura que pode não provocar resposta é:",
        "options": {
            "A": "±0,02 °C.",
            "B": "±0,10 °C.",
            "C": "±0,20 °C.",
            "D": "±1,00 °C.",
            "E": "±2,00 °C.",
        },
        "answer": "C",
    },
    {
        "text": "Um voltímetro analógico possui fundo de escala de 300 V e classe de exatidão de 1,5% do FS. O erro máximo absoluto especificado é:",
        "options": {
            "A": "±1,5 V.",
            "B": "±3,0 V.",
            "C": "±4,5 V.",
            "D": "±6,0 V.",
            "E": "±9,0 V.",
        },
        "answer": "C",
    },
    {
        "text": "Um transdutor de pressão opera de -14 a 236 psi e sua saída varia de 875 a 375 mV, em relação inversa. A sensibilidade aproximada é:",
        "options": {
            "A": "+0,5 mV/psi.",
            "B": "-0,5 mV/psi.",
            "C": "+2,0 mV/psi.",
            "D": "-2,0 mV/psi.",
            "E": "-4,0 mV/psi.",
        },
        "answer": "D",
    },
    {
        "text": "Na identificação ISA, a tag TIC-103 indica, em condições usuais, um instrumento associado a:",
        "options": {
            "A": "vazão, indicação e chave, malha 103.",
            "B": "temperatura, indicação e controle, malha 103.",
            "C": "pressão, transmissão e controle, malha 103.",
            "D": "nível, indicação e conversão, malha 103.",
            "E": "temperatura, alarme e registro, malha 103.",
        },
        "answer": "B",
    },
    {
        "text": "Segundo a lógica de identificação funcional da ISA, um transmissor de pressão diferencial usado para medir nível deve ser identificado como:",
        "options": {
            "A": "PDT, pois seu elemento mede diferença de pressão.",
            "B": "PT, pois a construção interna é de pressão.",
            "C": "LT, pois a função de processo é nível.",
            "D": "LI, pois todo nível precisa de indicação local.",
            "E": "LS, pois o transmissor atua como uma chave.",
        },
        "answer": "C",
    },
    {
        "text": "Um P&ID elaborado segundo a ISA S5.1 deve representar principalmente:",
        "options": {
            "A": "cotas civis e fundações da planta.",
            "B": "instrumentos, tags e linhas de sinal.",
            "C": "somente a sequência de produção.",
            "D": "listas de compras e custos totais.",
            "E": "escalas de desenho arquitetônico.",
        },
        "answer": "B",
    },
    {
        "text": "Um manômetro indica 10 psig. Considerando a pressão atmosférica de 14,7 psi, a pressão absoluta é:",
        "options": {
            "A": "4,7 psia.",
            "B": "10,0 psia.",
            "C": "14,7 psia.",
            "D": "24,7 psia.",
            "E": "147,0 psia.",
        },
        "answer": "D",
    },
    {
        "text": "No paradoxo hidrostático de Stevin, a pressão no fundo de recipientes com o mesmo líquido e mesma altura de coluna depende:",
        "options": {
            "A": "apenas da área do fundo.",
            "B": "da forma do recipiente.",
            "C": "do volume total contido.",
            "D": "da altura da coluna líquida.",
            "E": "da largura da superfície.",
        },
        "answer": "D",
    },
    {
        "text": "Sobre manômetros de coluna líquida, assinale a alternativa correta:",
        "options": {
            "A": "são indicados para transmissão remota.",
            "B": "dispensam fluido de referência interno.",
            "C": "servem bem como indicação local.",
            "D": "medem sempre pressão absoluta.",
            "E": "substituem todo transmissor digital.",
        },
        "answer": "C",
    },
    {
        "text": "Em transmissores industriais de pressão, o diafragma é muito empregado porque:",
        "options": {
            "A": "apresenta grande deslocamento angular.",
            "B": "converte pressão direto em código digital.",
            "C": "permite boa sensibilidade em pequenos deslocamentos.",
            "D": "elimina a necessidade de calibração de fábrica.",
            "E": "só funciona em pressões atmosféricas.",
        },
        "answer": "C",
    },
    {
        "text": "O termopar industrial se baseia no efeito Seebeck e, por isso:",
        "options": {
            "A": "mede resistência elétrica absoluta.",
            "B": "mede diferença de temperatura entre juntas.",
            "C": "dispensa compensação da junta fria.",
            "D": "usa sempre ponte de Wheatstone.",
            "E": "produz sinal padronizado de 4 a 20 mA.",
        },
        "answer": "B",
    },
    {
        "text": "O uso de poços térmicos em sensores intrusivos de temperatura tem como efeito típico:",
        "options": {
            "A": "reduzir o tempo de resposta do sensor.",
            "B": "eliminar perturbação ao escoamento.",
            "C": "permitir troca sem contato direto com o processo.",
            "D": "substituir qualquer transmissor de campo.",
            "E": "impedir totalmente erros de calibração.",
        },
        "answer": "C",
    },
    {
        "text": "Em uma planta industrial, deseja-se selecionar sensores de temperatura para duas aplicações distintas:\n\nI. Medição contínua da temperatura do ar em um sistema HVAC, com baixo custo e boa sensibilidade.\nII. Proteção térmica de um motor, gerando atuação quando a temperatura ultrapassar um valor crítico.\n\nConsiderando o comportamento dos termistores NTC e PTC, a escolha mais adequada é:",
        "options": {
            "A": "utilizar NTC em I e PTC em II, pois o NTC é sensível para medição contínua e o PTC pode atuar como proteção térmica.",
            "B": "utilizar PTC em I e NTC em II, pois o PTC apresenta resposta mais linear e o NTC atua como chave térmica.",
            "C": "utilizar NTC nas duas aplicações, pois sua resistência aumenta bruscamente em temperaturas críticas.",
            "D": "utilizar PTC nas duas aplicações, pois sua resistência diminui continuamente com o aumento da temperatura.",
            "E": "evitar termistores em I e II, pois eles não apresentam variação elétrica com a temperatura.",
        },
        "answer": "A",
    },
    {
        "text": "Um Pt100 é um RTD de platina que apresenta:",
        "options": {
            "A": "100 ohms a 0 °C e coeficiente positivo.",
            "B": "100 ohms a 100 °C e coeficiente positivo.",
            "C": "0 ohm a 100 °C e coeficiente negativo.",
            "D": "100 ohms a 0 °C e coeficiente negativo.",
            "E": "0 ohm a 100 °C e coeficiente positivo.",
        },
        "answer": "A",
    },
    {
        "text": "Na medição de nível por pressão diferencial em vaso fechado com perna seca:",
        "options": {
            "A": "a pressão de vapor age só no lado de alta.",
            "B": "a pressão de vapor se cancela nos dois lados.",
            "C": "o transmissor mede diretamente volume.",
            "D": "a densidade do líquido deixa de importar.",
            "E": "a pressão de vapor age só no lado de baixa.",
        },
        "answer": "B",
    },
    {
        "text": "Em relação aos medidores de nível por ultrassom e por radar, analise as afirmativas a seguir.\n\nI. O sensor ultrassônico utiliza ondas mecânicas; por isso, sua medição depende das propriedades do meio de propagação.\nII. O radar utiliza ondas eletromagnéticas na faixa de micro-ondas, podendo operar em atmosferas com vapor e até em vácuo.\nIII. Ambos podem ser instalados no topo do reservatório e realizar medição sem contato direto com o produto.\nIV. O ultrassom é mais adequado que o radar para ambientes com vácuo, pois sua onda não precisa de meio material.\n\nAssinale a alternativa correta.",
        "options": {
            "A": "Apenas I e II estão corretas.",
            "B": "Apenas II e IV estão corretas.",
            "C": "Apenas I, II e III estão corretas.",
            "D": "Apenas I, III e IV estão corretas.",
            "E": "Todas as afirmativas estão corretas.",
        },
        "answer": "C",
    },
]


INSTRUCTIONS = [
    "Assinale apenas uma alternativa correta em cada questão.",
    "Marque as respostas exclusivamente na folha de respostas.",
    "Questões ou marcações rasuradas no gabarito não serão aceitas.",
    "Use caneta azul ou preta para preencher a folha de respostas.",
    "Não é permitido o uso de celulares, relógios inteligentes, calculadoras, fones de ouvido ou quaisquer aparelhos eletroeletrônicos.",
    "Não é permitida consulta a materiais impressos, anotações ou comunicação com colegas.",
    "O tempo total de prova é de 100 minutos.",
    "Nenhum aluno poderá se ausentar da sala antes de decorridos 30 minutos do início da prova.",
    "Ao terminar, entregue a prova e a folha de respostas ao professor.",
]


def ensure_out() -> None:
    OUT.mkdir(exist_ok=True)


def grayscale_logo_path() -> Path:
    target = OUT / "cefet_mg_logo_pb.png"
    if LOGO.exists():
        logo = Image.open(LOGO).convert("RGBA")
        gray = ImageOps.grayscale(logo)
        alpha = logo.getchannel("A")
        bw_logo = Image.merge("RGBA", (gray, gray, gray, alpha))
        bw_logo.save(target)
        return target
    return LOGO


def make_versions() -> dict[str, list[dict[str, object]]]:
    versions: dict[str, list[dict[str, object]]] = {}
    for version_index, version in enumerate(VERSIONS):
        version_questions = []
        for question_index, question in enumerate(QUESTIONS):
            items = list(question["options"].items())
            rng = random.Random(RANDOM_SEED + version_index * 1000 + question_index)
            rng.shuffle(items)
            shuffled = {}
            answer = ""
            for option_index, (original_label, option_text) in enumerate(items):
                new_label = chr(ord("A") + option_index)
                shuffled[new_label] = option_text
                if original_label == question["answer"]:
                    answer = new_label
            version_questions.append(
                {
                    "number": question_index + 1,
                    "part": question.get("part"),
                    "text": question["text"],
                    "options": shuffled,
                    "answer": answer,
                }
            )
        versions[version] = version_questions
    return versions


def wrap_lines(text: str, max_width: float, font: str, size: float) -> list[str]:
    lines: list[str] = []
    for paragraph in text.split("\n"):
        words = paragraph.split()
        if not words:
            lines.append("")
            continue
        line = words[0]
        for word in words[1:]:
            candidate = f"{line} {word}"
            if stringWidth(candidate, font, size) <= max_width:
                line = candidate
            else:
                lines.append(line)
                line = word
        lines.append(line)
    return lines


def draw_wrapped(
    c: canvas.Canvas,
    text: str,
    x: float,
    y: float,
    max_width: float,
    font: str = "Helvetica",
    size: float = 10,
    leading: float = 13,
) -> float:
    c.setFont(font, size)
    for line in wrap_lines(text, max_width, font, size):
        if line:
            c.drawString(x, y, line)
        y -= leading
    return y


def question_height(question: dict[str, object], max_width: float) -> float:
    height = 18
    if question.get("part"):
        height += 32
    height += len(wrap_lines(str(question["text"]), max_width, "Helvetica", 10)) * 13
    height += 6
    for label, option in question["options"].items():
        height += len(wrap_lines(f"{label}) {option}", max_width - 16, "Helvetica", 9.5)) * 12
    return height + 12


def draw_part_heading(c: canvas.Canvas, text: str, x: float, y: float, max_width: float) -> float:
    c.setStrokeColorRGB(*DARK_GRAY)
    c.setLineWidth(0.8)
    c.roundRect(x, y - 20, max_width, 24, 3, fill=0, stroke=1)
    c.setFillColorRGB(*BLACK)
    c.setFont("Helvetica-Bold", 10)
    c.drawString(x + 12, y - 11, text)
    c.setFillColorRGB(*TEXT_GRAY)
    return y - 34


def draw_logo(c: canvas.Canvas, x: float, y: float, size: float) -> None:
    if LOGO.exists():
        try:
            logo = Image.open(LOGO).convert("RGBA")
            gray = ImageOps.grayscale(logo)
            alpha = logo.getchannel("A")
            bw_logo = Image.merge("RGBA", (gray, gray, gray, alpha))
            c.drawImage(ImageReader(bw_logo), x, y, width=size, height=size, mask="auto")
            return
        except Exception:
            pass
    c.setStrokeColorRGB(*BLACK)
    c.rect(x, y, size, size, fill=0, stroke=1)
    c.setFillColorRGB(*BLACK)
    c.setFont("Helvetica-Bold", 7)
    c.drawCentredString(x + size / 2, y + size / 2 - 2, "CEFET")


def footer(c: canvas.Canvas, page_no: int) -> None:
    c.setStrokeColorRGB(0.75, 0.75, 0.75)
    c.line(36, 36, PAGE_W - 36, 36)
    c.setFillColorRGB(0.35, 0.35, 0.35)
    c.setFont("Helvetica", 8)
    c.drawString(36, 24, FOOTER_TITLE)
    c.drawRightString(PAGE_W - 36, 24, f"Página {page_no}")


def draw_exam_header(c: canvas.Canvas, version: str, first_page: bool) -> float:
    c.setStrokeColorRGB(*BLACK)
    c.setLineWidth(1.2)
    c.line(36, PAGE_H - 28, PAGE_W - 36, PAGE_H - 28)
    c.setStrokeColorRGB(*MID_GRAY)
    c.setLineWidth(0.7)
    c.line(36, PAGE_H - 86, PAGE_W - 36, PAGE_H - 86)
    draw_logo(c, 36, PAGE_H - 82, 46)

    c.setFillColorRGB(*TEXT_GRAY)
    c.setFont("Helvetica-Bold", 10)
    c.drawString(92, PAGE_H - 44, "CENTRO FEDERAL DE EDUCAÇÃO TECNOLÓGICA DE MINAS GERAIS")
    c.setFont("Helvetica-Bold", 11)
    c.drawString(92, PAGE_H - 61, EXAM_TITLE)
    c.setFont("Helvetica", 9)
    c.drawString(92, PAGE_H - 76, COURSE_LINE)

    c.setStrokeColorRGB(*BLACK)
    c.setLineWidth(0.9)
    c.roundRect(PAGE_W - 106, PAGE_H - 76, 70, 34, 4, fill=0, stroke=1)
    c.setFillColorRGB(*BLACK)
    c.setFont("Helvetica-Bold", 9)
    c.drawCentredString(PAGE_W - 71, PAGE_H - 55, f"TIPO {version}")

    y = PAGE_H - 104
    c.setFillColorRGB(*TEXT_GRAY)
    c.setFont("Helvetica", 9)
    c.drawString(36, y, "Nome:")
    c.line(68, y - 2, 280, y - 2)
    c.drawString(292, y, "Turma:")
    c.line(326, y - 2, 405, y - 2)
    c.drawString(418, y, "Data:")
    c.line(447, y - 2, PAGE_W - 36, y - 2)
    y -= 16
    c.drawString(36, y, f"Valor: {TOTAL_POINTS:.1f} pontos | 20 questões | {POINTS_PER_QUESTION:.2f} ponto por questão | Duração: 100 minutos")

    if not first_page:
        return y - 24

    y -= 18
    c.setFillColorRGB(*LIGHT_GRAY)
    c.setStrokeColorRGB(*MID_GRAY)
    c.setLineWidth(0.6)
    c.roundRect(36, y - 104, PAGE_W - 72, 112, 4, fill=1, stroke=1)
    c.setFillColorRGB(*BLACK)
    c.setFont("Helvetica-Bold", 9)
    c.drawString(48, y - 8, "Instruções")
    c.setFillColorRGB(*TEXT_GRAY)
    c.setFont("Helvetica", 7.8)
    bullet_y = y - 22
    for item in INSTRUCTIONS:
        c.drawString(50, bullet_y, "-")
        bullet_y = draw_wrapped(c, item, 60, bullet_y, PAGE_W - 112, "Helvetica", 7.8, 10)
    return y - 126


def generate_exam_pdfs(versions: dict[str, list[dict[str, object]]]) -> None:
    for version, questions in versions.items():
        path = OUT / f"prova_sistemas_controle_tipo_{version}.pdf"
        c = canvas.Canvas(str(path), pagesize=A4)
        page_no = 1
        y = draw_exam_header(c, version, first_page=True)
        bottom = 58
        x = 42
        max_width = PAGE_W - 84

        for question in questions:
            needed = question_height(question, max_width)
            if y - needed < bottom:
                footer(c, page_no)
                c.showPage()
                page_no += 1
                y = draw_exam_header(c, version, first_page=False)

            if question.get("part"):
                y = draw_part_heading(c, str(question["part"]), x, y, max_width)

            c.setFillColorRGB(*BLACK)
            c.setFont("Helvetica-Bold", 10)
            c.drawString(x, y, f"{question['number']}.")
            c.setFillColorRGB(*TEXT_GRAY)
            y = draw_wrapped(c, str(question["text"]), x + 20, y, max_width - 20, "Helvetica", 10, 13)
            y -= 3
            for label, option in question["options"].items():
                y = draw_wrapped(c, f"{label}) {option}", x + 20, y, max_width - 20, "Helvetica", 9.5, 12)
            y -= 8

        footer(c, page_no)
        c.save()


def add_docx_header(doc: Document, version: str) -> None:
    logo_path = grayscale_logo_path()
    if logo_path.exists():
        paragraph = doc.add_paragraph()
        run = paragraph.add_run()
        run.add_picture(str(logo_path), width=Cm(1.6))
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("CENTRO FEDERAL DE EDUCAÇÃO TECNOLÓGICA DE MINAS GERAIS\n")
    run.bold = True
    run.font.size = Pt(10)
    run = title.add_run(f"{EXAM_TITLE}\n")
    run.bold = True
    run.font.size = Pt(12)
    run = title.add_run(COURSE_LINE)
    run.font.size = Pt(10)
    doc.add_paragraph(f"Nome: ____________________________________  Turma: __________  Data: ____/____/______  Tipo: {version}")
    doc.add_paragraph(f"Valor: {TOTAL_POINTS:.1f} pontos | 20 questões | {POINTS_PER_QUESTION:.2f} ponto por questão | Duração: 100 minutos")
    p = doc.add_paragraph()
    p.add_run("Instruções").bold = True
    for item in INSTRUCTIONS:
        doc.add_paragraph(item, style="List Bullet")


def generate_docx(versions: dict[str, list[dict[str, object]]]) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)

    for version_index, (version, questions) in enumerate(versions.items()):
        if version_index:
            doc.add_page_break()
        add_docx_header(doc, version)
        for question in questions:
            if question.get("part"):
                p = doc.add_paragraph()
                run = p.add_run(str(question["part"]))
                run.bold = True
                run.font.size = Pt(12)
            p = doc.add_paragraph()
            p.add_run(f"{question['number']}. ").bold = True
            p.add_run(str(question["text"]))
            for label, option in question["options"].items():
                doc.add_paragraph(f"{label}) {option}", style=None)
    doc.save(OUT / "prova_sistemas_controle_versoes_A_B_C_D.docx")


def answer_sheet_layout() -> dict[str, object]:
    box_left = 42.0
    box_right = PAGE_W - 42.0
    box_top = 620.0
    box_bottom = 120.0
    marker_margin = 16.0
    marker_size = 18.0

    markers = {
        "tl": [box_left + marker_margin, box_top - marker_margin],
        "tr": [box_right - marker_margin, box_top - marker_margin],
        "bl": [box_left + marker_margin, box_bottom + marker_margin],
        "br": [box_right - marker_margin, box_bottom + marker_margin],
    }
    version_y = 570.0
    version_bubbles = {
        "A": [225.0, version_y],
        "B": [280.0, version_y],
        "C": [335.0, version_y],
        "D": [390.0, version_y],
    }
    answer_bubbles: dict[str, dict[str, list[float]]] = {}
    x_positions = {"A": 185.0, "B": 240.0, "C": 295.0, "D": 350.0, "E": 405.0}
    start_y = 526.0
    row_gap = 18.0
    for question in range(1, 21):
        row_y = start_y - (question - 1) * row_gap
        answer_bubbles[str(question)] = {label: [x, row_y] for label, x in x_positions.items()}

    return {
        "page_width_pt": PAGE_W,
        "page_height_pt": PAGE_H,
        "marker_size_pt": marker_size,
        "bubble_radius_pt": 7.0,
        "markers_pt": markers,
        "version_bubbles_pt": version_bubbles,
        "answer_bubbles_pt": answer_bubbles,
    }


def draw_bubble(c: canvas.Canvas, x: float, y: float, radius: float = 7.0) -> None:
    c.setStrokeColorRGB(0, 0, 0)
    c.setLineWidth(0.8)
    c.circle(x, y, radius, stroke=1, fill=0)


def draw_marker(c: canvas.Canvas, center: list[float], size: float) -> None:
    x, y = center
    c.setFillColorRGB(0, 0, 0)
    c.rect(x - size / 2, y - size / 2, size, size, fill=1, stroke=0)


def generate_answer_sheet() -> dict[str, object]:
    layout = answer_sheet_layout()
    path = OUT / "folha_respostas_20_questoes.pdf"
    c = canvas.Canvas(str(path), pagesize=A4)

    c.setStrokeColorRGB(*BLACK)
    c.setLineWidth(1.2)
    c.line(36, PAGE_H - 28, PAGE_W - 36, PAGE_H - 28)
    c.setStrokeColorRGB(*MID_GRAY)
    c.setLineWidth(0.7)
    c.line(36, PAGE_H - 86, PAGE_W - 36, PAGE_H - 86)
    draw_logo(c, 36, PAGE_H - 82, 46)
    c.setFillColorRGB(*TEXT_GRAY)
    c.setFont("Helvetica-Bold", 10)
    c.drawString(92, PAGE_H - 44, "CENTRO FEDERAL DE EDUCAÇÃO TECNOLÓGICA DE MINAS GERAIS")
    c.setFont("Helvetica-Bold", 11)
    c.drawString(92, PAGE_H - 63, ANSWER_SHEET_TITLE)
    c.setFont("Helvetica", 9)
    c.drawString(92, PAGE_H - 78, f"{COURSE_LINE} | Valor: {TOTAL_POINTS:.1f} pontos | 20 questões")

    y = PAGE_H - 112
    c.setFont("Helvetica", 10)
    c.drawString(42, y, "Nome:")
    c.line(78, y - 2, PAGE_W - 42, y - 2)
    y -= 28
    c.drawString(42, y, "Turma:")
    c.line(80, y - 2, 240, y - 2)
    c.drawString(270, y, "Data:")
    c.line(302, y - 2, 430, y - 2)
    c.drawString(455, y, "Nota:")
    c.line(486, y - 2, PAGE_W - 42, y - 2)

    c.setStrokeColorRGB(*MID_GRAY)
    c.setLineWidth(0.8)
    c.roundRect(42, 120, PAGE_W - 84, 500, 4, fill=0, stroke=1)
    for center in layout["markers_pt"].values():
        draw_marker(c, center, layout["marker_size_pt"])

    c.setFillColorRGB(*BLACK)
    c.setFont("Helvetica-Bold", 11)
    c.drawCentredString(PAGE_W / 2, 595, "Marque o tipo de prova")
    for label, center in layout["version_bubbles_pt"].items():
        c.setFillColorRGB(*TEXT_GRAY)
        c.setFont("Helvetica-Bold", 9)
        c.drawCentredString(center[0], center[1] + 15, label)
        draw_bubble(c, center[0], center[1], 8.0)

    c.setFillColorRGB(*TEXT_GRAY)
    c.setFont("Helvetica", 8.5)
    c.drawCentredString(PAGE_W / 2, 548, "Preencha completamente a bolha escolhida. Marcações rasuradas ou duplas não serão aceitas.")

    c.setFont("Helvetica-Bold", 9)
    for label, x in {"A": 185.0, "B": 240.0, "C": 295.0, "D": 350.0, "E": 405.0}.items():
        c.drawCentredString(x, 538, label)

    c.setFont("Helvetica", 9)
    for question, bubbles in layout["answer_bubbles_pt"].items():
        y = bubbles["A"][1]
        c.setFillColorRGB(*TEXT_GRAY)
        c.drawRightString(118, y - 3, f"{int(question):02d}")
        for label, center in bubbles.items():
            draw_bubble(c, center[0], center[1])

    c.setFont("Helvetica", 8)
    c.setFillColorRGB(0.35, 0.35, 0.35)
    c.drawString(42, 92, "Para corrigir por câmera, fotografe a folha inteira, com os quatro marcadores pretos visíveis.")
    c.drawString(42, 78, "Não dobre, não amasse e não escreva dentro da área cinza além das marcações das respostas.")
    c.save()
    return layout


def generate_workbook(versions: dict[str, list[dict[str, object]]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Gabaritos"
    headers = ["Tipo"] + [f"Q{i}" for i in range(1, 21)] + ["Valor por questão", "Total"]
    ws.append(headers)
    answer_keys: dict[str, list[str]] = {}
    for version, questions in versions.items():
        answers = [str(q["answer"]) for q in questions]
        answer_keys[version] = answers
        ws.append([version] + answers + [POINTS_PER_QUESTION, TOTAL_POINTS])

    ws2 = wb.create_sheet("Resultados")
    ws2.append(["Data/hora", "Nome", "Turma", "Tipo", "Acertos", "Nota"] + [f"Q{i}" for i in range(1, 21)] + ["Avisos"])
    wb.save(OUT / "gabaritos_e_resultados.xlsx")

    payload = {
        "total_points": TOTAL_POINTS,
        "points_per_question": POINTS_PER_QUESTION,
        "question_count": len(QUESTIONS),
        "versions": answer_keys,
    }
    (OUT / "gabaritos.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def generate_layout_json(layout: dict[str, object]) -> None:
    (OUT / "layout_gabarito.json").write_text(json.dumps(layout, ensure_ascii=False, indent=2), encoding="utf-8")


def generate_readme() -> None:
    text = f"""# {EXAM_TITLE}

Arquivos gerados para aplicação e correção da prova objetiva.
O layout de impressão foi adaptado para preto e branco, com hierarquia por linhas, bordas e contraste em escala de cinza.

## Arquivos principais

- `prova_sistemas_controle_tipo_A.pdf` a `prova_sistemas_controle_tipo_D.pdf`: quatro versões para impressão.
- `prova_sistemas_controle_versoes_A_B_C_D.docx`: versão editável com as quatro provas.
- `folha_respostas_20_questoes.pdf`: folha única de respostas com tipo A/B/C/D.
- `gabaritos_e_resultados.xlsx`: gabaritos por versão e modelo de resultados.
- `gabaritos.json` e `layout_gabarito.json`: arquivos usados pelo corretor local.

## Aplicação

- Valor total: {TOTAL_POINTS:.1f} pontos.
- Total de questões: 20.
- Valor por questão: {POINTS_PER_QUESTION:.2f} ponto.
- Tempo de prova: 100 minutos.
- O aluno não poderá se ausentar antes de 30 minutos.
- Marcações rasuradas, duplas ou fora das bolhas não devem ser aceitas.

## Correção por câmera

Execute `corretor_gabaritos.py` e acesse o endereço informado no celular. Use a câmera do celular para fotografar a folha inteira, mantendo os quatro marcadores pretos visíveis.

## Identidade visual

O cabeçalho foi montado com referência na página oficial de identidade visual da SECOM CEFET-MG:
https://www.secom.cefetmg.br/identidade-visual-do-cefet-mg/

A imagem local `cefet_mg_logo.png` foi obtida de:
https://commons.wikimedia.org/wiki/File:Logo_CEFET-MG.png
"""
    (OUT / "README.txt").write_text(text, encoding="utf-8")


def main() -> None:
    ensure_out()
    versions = make_versions()
    generate_exam_pdfs(versions)
    generate_docx(versions)
    layout = generate_answer_sheet()
    generate_workbook(versions)
    generate_layout_json(layout)
    generate_readme()
    print(f"Arquivos gerados em: {OUT}")
    for version in VERSIONS:
        answers = "".join(q["answer"] for q in versions[version])
        print(f"Tipo {version}: {answers}")


if __name__ == "__main__":
    main()
